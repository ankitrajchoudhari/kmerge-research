"""Build K_Merge_Phase3_Report.docx from K_Merge_Phase3_Report.md (needs pandoc and python-docx).

1. reference.docx: pandoc's default styles with Calibri/Cambria, deck colours, A4 page.
2. pandoc markdown -> docx (TeX math becomes native Word equations).
3. Post-process tables: DXA widths (not percentages), light borders, shaded header row, 9 pt text.
"""
import os
import subprocess
import tempfile
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
SRC_MD = f"{HERE}/K_Merge_Phase3_Report.md"
REF = os.path.join(tempfile.gettempdir(), "kmerge_reference.docx")
OUT = f"{HERE}/K_Merge_Phase3_Report.docx"

INK = RGBColor(0x1B, 0x26, 0x3B)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GREY = RGBColor(0x4A, 0x55, 0x63)

PAGE_W, PAGE_H = 11906, 16838          # A4 in twentieths of a point
MARGIN = 1247                          # 2.2 cm
TEXT_W = PAGE_W - 2 * MARGIN           # 9412 twips


def make_reference():
    default = os.path.join(tempfile.gettempdir(), "kmerge_ref_default.docx")
    subprocess.run(["pandoc", "-o", default,
                    "--print-default-data-file", "reference.docx"], check=True)
    d = Document(default)
    st = d.styles

    by_name = {x.name: x for x in st}

    def font(style, name, size=None, color=None, bold=None):
        s = by_name[style]
        s.font.name = name
        rpr = s.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(a), name)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(a)) is not None:
                del rfonts.attrib[qn(a)]
        if size:
            s.font.size = Pt(size)
        if color is not None:
            s.font.color.rgb = color
        if bold is not None:
            s.font.bold = bold

    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text",
                 "Footnote Text", "Caption", "Image Caption", "Table Caption"):
        font(name, "Calibri", 11 if name not in ("Compact", "Caption", "Image Caption",
                                                 "Table Caption", "Footnote Text") else None,
             INK)
    font("Compact", "Calibri", 10, INK)
    font("Image Caption", "Calibri", 9, GREY)
    font("Title", "Cambria", 24, INK, True)
    font("Heading 1", "Cambria", 22, INK, True)
    font("Heading 2", "Cambria", 15, TEAL, True)
    font("Heading 3", "Cambria", 12.5, INK, True)
    for name in ("Body Text", "First Paragraph"):
        pf = by_name[name].paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(7)
        pf.line_spacing = 1.08
    by_name["Heading 2"].paragraph_format.space_before = Pt(18)
    by_name["Heading 2"].paragraph_format.space_after = Pt(6)
    by_name["Heading 2"].paragraph_format.keep_with_next = True
    by_name["Heading 1"].paragraph_format.space_after = Pt(4)
    by_name["Image Caption"].paragraph_format.space_after = Pt(12)

    sec = d.sections[0]
    sec.page_width, sec.page_height = Emu(PAGE_W * 635), Emu(PAGE_H * 635)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Emu(MARGIN * 635))
    sec.top_margin = sec.bottom_margin = Emu(1134 * 635)
    d.save(REF)


def run_pandoc():
    subprocess.run(["pandoc", SRC_MD, "-f", "markdown", "-t", "docx",
                    f"--reference-doc={REF}", f"--resource-path={HERE}", "-o", OUT],
                   check=True, cwd=HERE)


def set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "C9D2DC")
        b.append(e)
    tblPr.append(b)
    mar = tblPr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = OxmlElement("w:tblCellMar")
        for edge, w in (("top", 50), ("left", 90), ("bottom", 50), ("right", 90)):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:w"), str(w))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tblPr.append(mar)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def fix_widths(tbl):
    grid = tbl._tbl.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    raw = [int(c.get(qn("w:w"))) for c in cols]
    scale = TEXT_W / sum(raw)
    widths = [round(w * scale) for w in raw]
    widths[-1] += TEXT_W - sum(widths)
    for c, w in zip(cols, widths):
        c.set(qn("w:w"), str(w))
    tblPr = tbl._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(TEXT_W))
    for row in tbl.rows:
        for cell, w in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.insert(0, tcW)
            tcW.set(qn("w:type"), "dxa")
            tcW.set(qn("w:w"), str(w))


def post_process():
    d = Document(OUT)
    for tbl in d.tables:
        fix_widths(tbl)
        set_borders(tbl)
        keep_together = len(tbl.rows) <= 8
        for r, row in enumerate(tbl.rows):
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                cs = OxmlElement("w:cantSplit")
                trPr.insert(0, cs)
            if keep_together and r < len(tbl.rows) - 1:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
        prev = tbl._tbl.getprevious()
        if prev is not None and prev.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            Paragraph(prev, d._body).paragraph_format.keep_with_next = True
        for r, row in enumerate(tbl.rows):
            for cell in row.cells:
                if r == 0:
                    shade(cell, "E4ECF1")
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if r == 0:
                            run.font.bold = True
    # keep each chart with its caption
    for p in d.paragraphs:
        if p.style.name in ("Captioned Figure", "Figure"):
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    core = d.core_properties
    core.title = "K-Merge Phase 3: Share the Slot, Don't Average It"
    core.author = "Ankit Raj Choudhary"
    core.subject = "Report for Prof. Divya Saxena, IIT Jodhpur (SAIDE)"
    d.save(OUT)


if __name__ == "__main__":
    make_reference()
    run_pandoc()
    post_process()
    subprocess.run([sys.executable, os.path.join(HERE, "schema_fix.py"), OUT], check=True)
    print("built", OUT)
