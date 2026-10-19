"""K-Merge Phase 3 'work left for full proof' tracker (openpyxl, Arial, formulas for every total)."""
import math
import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

AUTHOR = "Ankit Raj Choudhary"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "K_Merge_Proof_Tracker.xlsx")

NAVY = "1B263B"
F = "Arial"
thin = Side(style="thin", color="C9D2DC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
FILL_IN = PatternFill("solid", fgColor="FFF2CC")      # light yellow: cells the user fills in
PART_FILL = {"1. Full proof": "E3F1F2", "2. Open items": "F1F3F5", "3. Own method": "EFEAF5", "4. Paper": "F7F1E8"}
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER_TOP = Alignment(wrap_text=True, vertical="top", horizontal="center")


def font(**kw):
    base = dict(name=F, size=10)
    base.update(kw)
    return Font(**base)


def set_height(ws, row, cols_widths, texts, line_h=13.5, min_lines=1):
    lines = min_lines
    for w, t in zip(cols_widths, texts):
        if t is None:
            continue
        chars = max(1, int(w * 1.0))
        n = sum(max(1, math.ceil(len(part) / chars)) for part in str(t).split("\n"))
        lines = max(lines, n)
    ws.row_dimensions[row].height = line_h * lines + 4


# ------------------------------------------------------------------ data (from the plan answer)
ITEMS = [
    # #, part, item, what it proves, planned GPU h, GPU as stated, status, when, pass rule
    (1, "1. Full proof", "Crowded-slot test, seeds 1 and 2",
     "K-Share beating split when all 12 tasks share a 4-direction slot (94.4% vs 87.1%) wasn't luck from one seed",
     5, "~5 h", "Ready to run", "Week 1",
     "K-Share beats split at 4 directions on all 3 seeds"),
    (2, "1. Full proof", "Real benchmark, seed 2",
     "The main result (split 98.9% vs K-Merge 90.8%) holds on 3 seeds, not just 2",
     4.5, "~4.5 h", "Ready to run", "Week 1",
     "Split beats K-Merge and RegMean with p < 0.05 on 3 seeds, at K = 2 and K = 4"),
    (3, "1. Full proof", "K-Share with exactly equal storage, real data",
     "K-Share's win doesn't come from its 2–8% extra storage",
     1, "under 1 h", "Small change", "Week 1",
     "Equal-storage K-Share stays within about 1 point of normal K-Share"),
    (4, "1. Full proof", "KnOTS baseline, real data",
     "Sharing also beats the strongest merging method reviewers will ask about",
     1, "under 1 h", "Code exists in the lab", "Week 2",
     "Sharing also beats KnOTS on the real tasks"),
    (5, "1. Full proof", "Strictly online sharing, real data",
     "Sharing still works when the phone keeps only the slot and never the original adapters, which is K-Merge's real setting. So far this is only tested in the lab.",
     1, "~1 h (my estimate)", "Lab code exists; needs moving to the real benchmark", "Week 1",
     "Online sharing stays within about 1–2 points of the rebuilt version"),
    (6, "1. Full proof", "Llama-3.2-1B, the K-Merge paper's other model",
     "The result isn't specific to Qwen",
     5, "~5 h", "Small code change", "Week 2",
     "Same order on Llama: sharing first, then the best merge, then K-Merge"),
    (7, "2. Open items", "Analysis 4 on 3 seeds",
     "The French recovery (0.860 → 0.944) is still from seed 0 only",
     1, "~1 h", "Not started", "Week 2", "—"),
    (8, "2. Open items", "Phase 1 checks (professor's items 20–21)",
     "The seed-0 DE-EN rows and the per-seed similarity scores",
     0, "no GPU", "Not started", "Now", "—"),
    (9, "2. Open items", "English re-score with official ERRANT (optional)",
     "Uses the saved outputs; answers \"your grammar metric only approximates ERRANT\"",
     0, "no GPU", "Optional; not started", "Now", "—"),
    (10, "3. Own method", "Rebuild the K-Share+ code, then GPU pilot",
     "The draft was lost when the old workspace reset. Test it without GPU first (start now), then run the GPU pilot. This could become our own method.",
     1, "~1 h", "Draft lost; needs rebuilding", "Week 1", "—"),
    (11, "4. Paper", "Phone cost: memory and rebuild time",
     "Supports the on-device claim; measured on a computer or a phone",
     0, "no GPU", "Not started", "Weeks 3–4", "—"),
    (12, "4. Paper", "Real Hindi grammar data",
     "Blocked until the authors of a Hindi corpus share their data; an ask for the professor",
     0, "—", "Blocked: waiting for data", "Blocked", "—"),
    (13, "4. Paper", "Add results to the shared Google Doc",
     "The professor's request; needs edit access from her",
     0, "—", "Waiting for edit access", "Weeks 3–4", "—"),
    (14, "4. Paper", "Paper draft",
     "Write it after items 1–6",
     0, "—", "After items 1–6", "Weeks 3–4", "—"),
]

wb = Workbook()

# ================================================================== Sheet 1: Plan
ws = wb.active
ws.title = "Plan"
cols = [("#", 5), ("Part", 15), ("Item", 32), ("What it proves / why", 46), ("GPU hours (planned)", 11),
        ("GPU time as stated", 14), ("Status (5 Oct 2026)", 22), ("When", 11), ("Pass rule", 36),
        ("Actual GPU hours", 11), ("Result / notes", 36), ("Done?", 8)]
for i, (_, w) in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws["A1"] = "K-Merge Phase 3: work left for full proof"
ws["A1"].font = font(size=14, bold=True, color=NAVY)
ws["A2"] = ("How to use: fill in the yellow cells (Actual GPU hours, Result / notes, Done?). Blue numbers are planned "
            "estimates; \"under 1 h\" is counted as 1 h. Totals are on the Summary sheet, pass rules on the Pass rules sheet.")
ws["A2"].font = font(italic=True, color="4A5563")
ws["A3"] = ("Example of a filled row (example only, not a real result):  Actual GPU hours = 4.6  ·  "
            "Result / notes = \"Seed 2 done; split still ahead of K-Merge and RegMean (Holm p < 0.05)\"  ·  Done? = Yes")
ws["A3"].font = font(italic=True, color="4A5563")
for r in (2, 3):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(cols))
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[2].height = 28
ws.row_dimensions[3].height = 16

HDR = 5
for i, (name, _) in enumerate(cols, 1):
    c = ws.cell(HDR, i, name)
    c.font = font(bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if i in (1, 5, 10, 12) else "left")
    c.border = BORDER
ws.row_dimensions[HDR].height = 30

first = HDR + 1
for k, (n, part, item, why, gpu, gpu_txt, status, when, rule) in enumerate(ITEMS):
    r = first + k
    vals = [n, part, item, why, gpu, gpu_txt, status, when, rule, None, None, "No"]
    for i, v in enumerate(vals, 1):
        c = ws.cell(r, i, v)
        c.font = font()
        c.border = BORDER
        c.alignment = CENTER_TOP if i in (1, 5, 8, 10, 12) else WRAP_TOP
    ws.cell(r, 2).fill = PatternFill("solid", fgColor=PART_FILL[part])
    ws.cell(r, 3).font = font(bold=True)
    ws.cell(r, 5).font = font(color="0000FF")          # hardcoded planned estimate
    ws.cell(r, 5).number_format = '0.0;-0.0;"–"'
    for i in (10, 11, 12):
        ws.cell(r, i).fill = FILL_IN
    ws.cell(r, 10).number_format = "0.0"
    set_height(ws, r, [w for _, w in cols], vals)
last = first + len(ITEMS) - 1

ws.cell(first, 5).comment = Comment("Planned GPU hours from the Phase 3 report's 'What is left' table; "
                                    "item 5 is an estimate. 'under 1 h' is entered as 1 h (upper bound).", AUTHOR)
ws.freeze_panes = ws.cell(first, 4)

dv_done = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(dv_done)
dv_done.add(f"L{first}:L{last}")
ws.conditional_formatting.add(f"L{first}:L{last}",
                              CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="CDEBD6")))
ws.conditional_formatting.add(f"A{first}:K{last}",
                              FormulaRule(formula=[f'$L{first}="Yes"'], font=Font(name=F, color="6B7785", strike=True)))
ws.auto_filter.ref = f"A{HDR}:L{last}"
ws.sheet_view.zoomScale = 100

# ================================================================== Sheet 2: Summary
sm = wb.create_sheet("Summary")
sm.column_dimensions["A"].width = 46
sm.column_dimensions["B"].width = 20
sm.column_dimensions["C"].width = 12
sm["A1"] = "Summary (updates as you fill in the Plan sheet)"
sm["A1"].font = font(size=14, bold=True, color=NAVY)

E = f"Plan!$E${first}:$E${last}"
B = f"Plan!$B${first}:$B${last}"
C = f"Plan!$C${first}:$C${last}"
H = f"Plan!$H${first}:$H${last}"
J = f"Plan!$J${first}:$J${last}"
L = f"Plan!$L${first}:$L${last}"

rows = [
    ("Planned GPU hours, all items", f"=SUM({E})", "0.0"),
    ("Planned GPU hours, Part 1 (full proof, items 1–6)", f'=SUMIF({B},"1. Full proof",{E})', "0.0"),
    ("Planned GPU hours still to do (Done? is not Yes)", f'=SUMIF({L},"<>Yes",{E})', "0.0"),
    ("Actual GPU hours used so far", f"=SUM({J})", "0.0"),
    ("Items done", f'=COUNTIF({L},"Yes")', "0"),
    ("Items in total", f"=COUNTA({C})", "0"),
    ("Kaggle free GPU quota per week (hours)", 30, "0"),
    ("Weeks of quota needed for the GPU work still to do", "=IF(B9=0,0,B5/B9)", "0.00"),
]
for k, (label, val, fmt) in enumerate(rows):
    r = 3 + k
    a = sm.cell(r, 1, label)
    a.font = font()
    a.border = BORDER
    b = sm.cell(r, 2, val)
    b.number_format = fmt
    b.border = BORDER
    b.alignment = Alignment(horizontal="right")
    b.font = font(bold=True, color="0000FF") if isinstance(val, (int, float)) else font(bold=True)
sm["B9"].comment = Comment("Kaggle's free quota as stated in our report (30 GPU hours a week). Change it if your quota differs.", AUTHOR)
sm["A11"] = "Plan spreads the GPU work over two weeks to leave room for reruns."
sm["A11"].font = font(italic=True, color="4A5563")

sm["A13"] = "By week"
sm["A13"].font = font(bold=True, size=11, color=NAVY)
for i, name in enumerate(["When", "Planned GPU hours", "Items"], 1):
    c = sm.cell(14, i, name)
    c.font = font(bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.border = BORDER
weeks = ["Now", "Week 1", "Week 2", "Weeks 3–4", "Blocked"]
for k, wk in enumerate(weeks):
    r = 15 + k
    sm.cell(r, 1, wk).font = font()
    sm.cell(r, 2, f'=SUMIF({H},A{r},{E})').number_format = "0.0"
    sm.cell(r, 3, f'=COUNTIF({H},A{r})').number_format = "0"
    for i in (1, 2, 3):
        sm.cell(r, i).border = BORDER
        if i > 1:
            sm.cell(r, i).font = font()
tr = 15 + len(weeks)
sm.cell(tr, 1, "Total").font = font(bold=True)
sm.cell(tr, 2, f"=SUM(B15:B{tr - 1})").number_format = "0.0"
sm.cell(tr, 3, f"=SUM(C15:C{tr - 1})").number_format = "0"
for i in (1, 2, 3):
    sm.cell(tr, i).border = BORDER
    sm.cell(tr, i).font = font(bold=True)
sm.cell(tr + 1, 1, "Check: total by week equals the Plan total").font = font(italic=True, color="4A5563")
sm.cell(tr + 1, 2, f'=IF(ABS(B{tr}-B3)<0.001,"OK","Mismatch")').font = font(italic=True)

# ================================================================== Sheet 3: Pass rules
pr = wb.create_sheet("Pass rules")
pcols = [("#", 5), ("Claim to prove", 30), ("Pass rule (fixed before running)", 40), ("Experiment", 11),
         ("Where we are now", 44), ("Result (fill in)", 34), ("Passed?", 10)]
for i, (_, w) in enumerate(pcols, 1):
    pr.column_dimensions[get_column_letter(i)].width = w
pr["A1"] = "Pass rules: when a claim counts as proved"
pr["A1"].font = font(size=14, bold=True, color=NAVY)
pr["A2"] = ("Write the result next to each rule after the runs. If one rule fails, drop only that claim; the others "
            "stand on their own. \"Where we are now\" is from the Phase 3 report (5 October 2026).")
pr["A2"].font = font(italic=True, color="4A5563")
pr.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(pcols))
pr["A2"].alignment = Alignment(wrap_text=True, vertical="top")
pr.row_dimensions[2].height = 28

RULES = [
    (1, "Main result: sharing beats merging on real tasks",
     "Split beats K-Merge and RegMean with p < 0.05 on 3 seeds, at K = 2 and K = 4", "Item 2",
     "2 seeds: split 98.9%, RegMean 95.6%, K-Merge 90.8% (K = 2); all differences p < 0.001 (Holm)"),
    (2, "Not specific to Qwen",
     "On Llama-3.2-1B, same order: sharing first, then the best merge, then K-Merge", "Item 6",
     "Not run yet"),
    (3, "K-Share's win is not from extra storage",
     "Equal-storage K-Share stays within about 1 point of normal K-Share", "Item 3",
     "K-Share 98.9% using 7.9% extra storage (K = 2); split 98.9% at exactly equal storage"),
    (4, "Works in K-Merge's real (online) setting",
     "Online sharing stays within about 1–2 points of the rebuilt version", "Item 5",
     "Tested only in the synthetic lab; real-data slots were rebuilt from the original adapters"),
    (5, "Crowded slots need a shared basis",
     "K-Share beats split at 4 directions on all 3 seeds", "Item 1",
     "1 seed: K-Share 94.4% vs split 87.1% (+7.3 points, 95% CI +5.2 to +9.2)"),
    (6, "Beats the strongest merge reviewers will ask about",
     "Sharing also beats KnOTS on the real tasks", "Item 4",
     "KnOTS not run on real data yet"),
]
PH = 4
for i, (name, _) in enumerate(pcols, 1):
    c = pr.cell(PH, i, name)
    c.font = font(bold=True, color="FFFFFF")
    c.fill = HEAD_FILL
    c.border = BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if i in (1, 4, 7) else "left")
pr.row_dimensions[PH].height = 30
for k, (n, claim, rule, exp, now) in enumerate(RULES):
    r = PH + 1 + k
    vals = [n, claim, rule, exp, now, None, "Pending"]
    for i, v in enumerate(vals, 1):
        c = pr.cell(r, i, v)
        c.font = font(bold=(i == 2))
        c.border = BORDER
        c.alignment = CENTER_TOP if i in (1, 4, 7) else WRAP_TOP
    for i in (6, 7):
        pr.cell(r, i).fill = FILL_IN
    set_height(pr, r, [w for _, w in pcols], vals)
plast = PH + len(RULES)
dv_pass = DataValidation(type="list", formula1='"Yes,No,Pending"', allow_blank=True)
pr.add_data_validation(dv_pass)
dv_pass.add(f"G{PH + 1}:G{plast}")
pr.conditional_formatting.add(f"G{PH + 1}:G{plast}",
                              CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="CDEBD6")))
pr.conditional_formatting.add(f"G{PH + 1}:G{plast}",
                              CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor="F6D4D8")))
pr.cell(plast + 2, 1, "Rules passed").font = font(bold=True)
pr.cell(plast + 2, 3, f'=COUNTIF(G{PH + 1}:G{plast},"Yes")&" of "&COUNTA(B{PH + 1}:B{plast})').font = font(bold=True)
pr.freeze_panes = pr.cell(PH + 1, 3)

for sheet in (ws, sm, pr):
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

wb.save(OUT)
print("saved", OUT, "items", len(ITEMS), "rows", first, "-", last)
